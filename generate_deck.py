import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

def build_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Color Palette Definitions
    BG_DARK = RGBColor(11, 8, 19)          # #0B0813
    CARD_BG = RGBColor(20, 14, 38)         # #140E26
    CARD_HEADER_BG = RGBColor(28, 20, 52)  # #1C1434
    BORDER_PURPLE = RGBColor(139, 92, 246) # #8B5CF6
    BORDER_DIM = RGBColor(45, 33, 74)      # #2D214A

    TEXT_WHITE = RGBColor(255, 255, 255)
    TEXT_MUTED = RGBColor(165, 180, 252)   # #A5B4FC
    TEXT_DIM = RGBColor(148, 163, 184)     # #94A3B8

    ACCENT_PURPLE = RGBColor(139, 92, 246) # #8B5CF6
    ACCENT_CYAN = RGBColor(0, 223, 137)    # #00DF89
    ACCENT_BLUE = RGBColor(59, 130, 246)   # #3B82F6
    ACCENT_AMBER = RGBColor(245, 158, 11)  # #F59E0B

    def set_slide_background(slide):
        bg = slide.background
        fill = bg.fill
        fill.solid()
        fill.fore_color.rgb = BG_DARK

        # Top subtle accent bar
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.08))
        bar.fill.solid()
        bar.fill.fore_color.rgb = ACCENT_PURPLE
        bar.line.fill.background()

    def add_header(slide, title_text, category_text="AGENTGRID HACKATHON PRESENTATION"):
        # Category Badge
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.35), Inches(11.7), Inches(0.3))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.size = Pt(10)
        p_cat.font.bold = True
        p_cat.font.color.rgb = ACCENT_CYAN
        p_cat.font.name = "Segoe UI"

        # Title Text
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(11.7), Inches(0.6))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = Pt(22)
        p_title.font.bold = True
        p_title.font.color.rgb = TEXT_WHITE
        p_title.font.name = "Segoe UI"

    def add_card(slide, left, top, width, height, bg_color=CARD_BG, border_color=BORDER_DIM):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        if border_color:
            card.line.color.rgb = border_color
            card.line.width = Pt(1.5)
        else:
            card.line.fill.background()
        return card

    # =========================================================================
    # SLIDE 1 — TEAM DETAILS
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1)

    # Hero Card Left
    add_card(s1, Inches(0.8), Inches(0.8), Inches(7.2), Inches(5.8), bg_color=CARD_BG, border_color=ACCENT_PURPLE)

    # Subtitle / Team Name
    tb1_brand = s1.shapes.add_textbox(Inches(1.2), Inches(1.2), Inches(6.4), Inches(0.4))
    p1 = tb1_brand.text_frame.paragraphs[0]
    p1.text = "TEAM TECHSPARK PRESENTS"
    p1.font.size = Pt(11)
    p1.font.bold = True
    p1.font.color.rgb = ACCENT_CYAN

    # Title
    tb1_title = s1.shapes.add_textbox(Inches(1.2), Inches(1.6), Inches(6.4), Inches(1.2))
    p2 = tb1_title.text_frame.paragraphs[0]
    p2.text = "AgentGrid"
    p2.font.size = Pt(48)
    p2.font.bold = True
    p2.font.color.rgb = TEXT_WHITE

    # Tagline
    tb1_tag = s1.shapes.add_textbox(Inches(1.2), Inches(2.7), Inches(6.4), Inches(0.5))
    p3 = tb1_tag.text_frame.paragraphs[0]
    p3.text = '"Where AI Agents Work Together."'
    p3.font.size = Pt(18)
    p3.font.italic = True
    p3.font.color.rgb = ACCENT_PURPLE

    # Concept summary in Hero Card
    tb1_desc = s1.shapes.add_textbox(Inches(1.2), Inches(3.4), Inches(6.4), Inches(1.5))
    tf1_desc = tb1_desc.text_frame
    tf1_desc.word_wrap = True
    p4 = tf1_desc.paragraphs[0]
    p4.text = "An autonomous multi-agent AI workforce for business execution. Translating high-level founder objectives into coordinated, dependency-aware workflows across Finance, Marketing, Hiring, and Legal departments."
    p4.font.size = Pt(13)
    p4.font.color.rgb = TEXT_MUTED

    # Team Roster Card Right
    add_card(s1, Inches(8.3), Inches(0.8), Inches(4.2), Inches(5.8), bg_color=CARD_BG, border_color=BORDER_DIM)

    tb1_team_hdr = s1.shapes.add_textbox(Inches(8.6), Inches(1.2), Inches(3.6), Inches(0.4))
    p_th = tb1_team_hdr.text_frame.paragraphs[0]
    p_th.text = "TEAM MEMBERS"
    p_th.font.size = Pt(12)
    p_th.font.bold = True
    p_th.font.color.rgb = ACCENT_CYAN

    members = [
        ("Yogender Verma", "Full-Stack & Multi-Agent Lead"),
        ("G. Sri Pranay", "AI Systems & Backend Architect"),
        ("Ch. Sai Sindhuri Reddy", "Frontend & UI/UX Engineer")
    ]

    for idx, (m_name, m_role) in enumerate(members):
        top_pos = Inches(1.8 + idx * 1.5)
        add_card(s1, Inches(8.6), top_pos, Inches(3.6), Inches(1.2), bg_color=CARD_HEADER_BG, border_color=BORDER_PURPLE)
        
        tb_m = s1.shapes.add_textbox(Inches(8.8), top_pos + Inches(0.15), Inches(3.2), Inches(0.9))
        tf_m = tb_m.text_frame
        tf_m.word_wrap = True
        
        pm1 = tf_m.paragraphs[0]
        pm1.text = m_name
        pm1.font.size = Pt(15)
        pm1.font.bold = True
        pm1.font.color.rgb = TEXT_WHITE
        
        pm2 = tf_m.add_paragraph()
        pm2.text = m_role
        pm2.font.size = Pt(10)
        pm2.font.color.rgb = TEXT_MUTED

    # Footer
    tb_foot = s1.shapes.add_textbox(Inches(0.8), Inches(6.8), Inches(11.7), Inches(0.4))
    pf = tb_foot.text_frame.paragraphs[0]
    pf.text = "TechSpark • AgentGrid Hackathon Jury Deck"
    pf.font.size = Pt(10)
    pf.font.color.rgb = TEXT_DIM

    # =========================================================================
    # SLIDE 2 — PROBLEM STATEMENT
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2)
    add_header(s2, "The Problem: AI Can Answer — But Business Needs Execution")

    # Left Container: Problem Points Card
    add_card(s2, Inches(0.8), Inches(1.3), Inches(7.5), Inches(5.3), bg_color=CARD_BG, border_color=BORDER_DIM)

    problems = [
        ("Siloed Business Operations", "Founders manage hiring, finance, marketing, and legal simultaneously with fragmented attention."),
        ("Interconnected Dependencies", "A hiring decision directly impacts runway, requires marketing outreach, and legal contracts."),
        ("Task vs. Workflow Gap", "Existing AI tools answer one-off questions but fail to execute end-to-end multi-step workflows."),
        ("Manual Context Switching", "Moving data manually between AI chatbots, spreadsheets, and apps creates massive overhead."),
        ("Lack of Human Oversight", "Fully unconstrained AI causes rogue actions; important decisions demand human approval gates.")
    ]

    for idx, (p_title, p_desc) in enumerate(problems):
        top_y = Inches(1.5 + idx * 0.98)
        
        # Indicator dot shape
        dot = s2.shapes.add_shape(MSO_SHAPE.OVAL, Inches(1.1), top_y + Inches(0.08), Inches(0.15), Inches(0.15))
        dot.fill.solid()
        dot.fill.fore_color.rgb = ACCENT_PURPLE
        dot.line.fill.background()

        tb_prob = s2.shapes.add_textbox(Inches(1.35), top_y - Inches(0.05), Inches(6.7), Inches(0.9))
        tf_p = tb_prob.text_frame
        tf_p.word_wrap = True
        
        pp1 = tf_p.paragraphs[0]
        pp1.text = p_title
        pp1.font.size = Pt(13)
        pp1.font.bold = True
        pp1.font.color.rgb = TEXT_WHITE
        
        pp2 = tf_p.add_paragraph()
        pp2.text = p_desc
        pp2.font.size = Pt(11)
        pp2.font.color.rgb = TEXT_MUTED

    # Right Container: Question Callout Card
    add_card(s2, Inches(8.6), Inches(1.3), Inches(3.9), Inches(5.3), bg_color=CARD_HEADER_BG, border_color=ACCENT_PURPLE)

    tb_q_hdr = s2.shapes.add_textbox(Inches(8.9), Inches(1.6), Inches(3.3), Inches(0.4))
    pqh = tb_q_hdr.text_frame.paragraphs[0]
    pqh.text = "CORE CHALLENGE"
    pqh.font.size = Pt(11)
    pqh.font.bold = True
    pqh.font.color.rgb = ACCENT_AMBER

    tb_q_body = s2.shapes.add_textbox(Inches(8.9), Inches(2.2), Inches(3.3), Inches(3.8))
    tf_qb = tb_q_body.text_frame
    tf_qb.word_wrap = True
    pqb = tf_qb.paragraphs[0]
    pqb.text = '"How can one high-level business objective become coordinated, actionable work across multiple functions while keeping humans in control?"'
    pqb.font.size = Pt(17)
    pqb.font.bold = True
    pqb.font.color.rgb = ACCENT_CYAN

    # =========================================================================
    # SLIDE 3 — SOLUTION
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3)
    add_header(s3, "AgentGrid — From Objective to Execution")

    # Workflow Steps Row (Top)
    steps = [
        ("1. Founder Objective", "High-level goal", ACCENT_CYAN),
        ("2. CEO Agent", "Master Orchestrator", ACCENT_PURPLE),
        ("3. Master Plan (DAG)", "Task graph & deps", ACCENT_BLUE),
        ("4. Founder Approval", "Level A Gate", ACCENT_AMBER),
        ("5. Execution", "Depts execute", ACCENT_PURPLE),
        ("6. Control Gate", "Consequential action", ACCENT_CYAN)
    ]

    for idx, (st_title, st_desc, st_color) in enumerate(steps):
        left_x = Inches(0.8 + idx * 2.0)
        add_card(s3, left_x, Inches(1.3), Inches(1.85), Inches(1.2), bg_color=CARD_HEADER_BG, border_color=st_color)
        
        tb_st = s3.shapes.add_textbox(left_x + Inches(0.1), Inches(1.4), Inches(1.65), Inches(1.0))
        tf_st = tb_st.text_frame
        tf_st.word_wrap = True
        
        pst1 = tf_st.paragraphs[0]
        pst1.text = st_title
        pst1.font.size = Pt(11)
        pst1.font.bold = True
        pst1.font.color.rgb = TEXT_WHITE
        
        pst2 = tf_st.add_paragraph()
        pst2.text = st_desc
        pst2.font.size = Pt(9)
        pst2.font.color.rgb = TEXT_MUTED

    # Concrete Example Section Card (Middle)
    add_card(s3, Inches(0.8), Inches(2.7), Inches(11.7), Inches(3.2), bg_color=CARD_BG, border_color=BORDER_DIM)

    tb_ex_hdr = s3.shapes.add_textbox(Inches(1.1), Inches(2.9), Inches(11.1), Inches(0.4))
    pehdr = tb_ex_hdr.text_frame.paragraphs[0]
    pehdr.text = 'REAL EXECUTION EXAMPLE: "Hire 2 frontend interns under ₹10,000/month"'
    pehdr.font.size = Pt(13)
    pehdr.font.bold = True
    pehdr.font.color.rgb = ACCENT_CYAN

    dept_examples = [
        ("FINANCE AGENT", "Models stipend cap feasibility, runway impact, and approves ₹10k/mo budget ceiling.", ACCENT_BLUE),
        ("MARKETING AGENT", "Generates social recruitment campaign copy & one-click LinkedIn broadcast collateral.", ACCENT_PURPLE),
        ("HIRING AGENT", "Generates 20 MCQs + Python coding challenge, screens resumes, & scores candidates.", ACCENT_CYAN),
        ("LEGAL AGENT", "Drafts offer letters, NDAs, IP assignments, & compliance covenants for founder approval.", ACCENT_AMBER)
    ]

    for idx, (d_name, d_desc, d_col) in enumerate(dept_examples):
        left_d = Inches(1.1 + idx * 2.8)
        add_card(s3, left_d, Inches(3.4), Inches(2.6), Inches(2.2), bg_color=CARD_HEADER_BG, border_color=d_col)
        
        tb_d = s3.shapes.add_textbox(left_d + Inches(0.15), Inches(3.5), Inches(2.3), Inches(2.0))
        tf_d = tb_d.text_frame
        tf_d.word_wrap = True
        
        pd1 = tf_d.paragraphs[0]
        pd1.text = d_name
        pd1.font.size = Pt(11)
        pd1.font.bold = True
        pd1.font.color.rgb = d_col
        
        pd2 = tf_d.add_paragraph()
        pd2.text = d_desc
        pd2.font.size = Pt(10)
        pd2.font.color.rgb = TEXT_MUTED

    # Bottom Closing Highlight Box
    add_card(s3, Inches(0.8), Inches(6.1), Inches(11.7), Inches(0.8), bg_color=CARD_HEADER_BG, border_color=ACCENT_PURPLE)
    tb_sol_close = s3.shapes.add_textbox(Inches(1.0), Inches(6.25), Inches(11.3), Inches(0.5))
    psc = tb_sol_close.text_frame.paragraphs[0]
    psc.alignment = PP_ALIGN.CENTER
    psc.text = "AI executes the work. The founder remains in control."
    psc.font.size = Pt(16)
    psc.font.bold = True
    psc.font.color.rgb = TEXT_WHITE

    # =========================================================================
    # SLIDE 4 — TECHNOLOGY STACK
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4)
    add_header(s4, "Technology Stack")

    tech_cards = [
        ("FRONTEND", [
            ("React 19", "Core UI Component Library"),
            ("TypeScript", "Strict Type Safety"),
            ("Vite", "Next-Gen Frontend Tooling"),
            ("Tailwind CSS", "Custom Glassmorphic Theme"),
            ("Framer Motion", "Micro-animations & Transitions"),
            ("Lucide React", "Icon System")
        ], ACCENT_CYAN, Inches(0.8), Inches(1.3)),

        ("BACKEND", [
            ("Python 3.11+", "Core Server Environment"),
            ("FastAPI", "High-performance Async Web API"),
            ("Uvicorn", "ASGI Production Server"),
            ("SQLAlchemy", "ORM & Database Management"),
            ("Pydantic v2", "Data Validation & Schemas"),
            ("SQLite / Postgres", "Flexible Relational Storage")
        ], ACCENT_BLUE, Inches(6.7), Inches(1.3)),

        ("AI & ORCHESTRATION", [
            ("Google Gemini API", "LLM Reasoning Engine"),
            ("Multi-Agent Framework", "Autonomous Executive Team"),
            ("CEO Agent", "Master Plan & DAG Decomposition"),
            ("DAG Execution Engine", "Dependency Graph Control"),
            ("Domain Models", "Specialized Agent Runtimes")
        ], ACCENT_PURPLE, Inches(0.8), Inches(4.3)),

        ("SECURITY & EXECUTION", [
            ("JWT Authentication", "Secure Stateless Auth"),
            ("Bcrypt / Passlib", "Password Encryption"),
            ("Human-in-the-Loop", "Multi-Level Gate Approvals"),
            ("CodeExecutor Sandbox", "Multi-Language Test Case Execution")
        ], ACCENT_AMBER, Inches(6.7), Inches(4.3))
    ]

    for cat_title, items, cat_col, left_pos, top_pos in tech_cards:
        add_card(s4, left_pos, top_pos, Inches(5.83), Inches(2.7), bg_color=CARD_BG, border_color=BORDER_DIM)
        
        # Category Header
        tb_t_hdr = s4.shapes.add_textbox(left_pos + Inches(0.3), top_pos + Inches(0.2), Inches(5.2), Inches(0.4))
        pth = tb_t_hdr.text_frame.paragraphs[0]
        pth.text = cat_title
        pth.font.size = Pt(12)
        pth.font.bold = True
        pth.font.color.rgb = cat_col

        # Items Grid
        tb_t_body = s4.shapes.add_textbox(left_pos + Inches(0.3), top_pos + Inches(0.6), Inches(5.23), Inches(1.9))
        tf_tb = tb_t_body.text_frame
        tf_tb.word_wrap = True

        for idx, (tech_name, tech_desc) in enumerate(items):
            p = tf_tb.paragraphs[0] if idx == 0 else tf_tb.add_paragraph()
            p.text = f"• {tech_name} — {tech_desc}"
            p.font.size = Pt(10)
            p.font.color.rgb = TEXT_WHITE

    # =========================================================================
    # SLIDE 5 — INNOVATION
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5)
    add_header(s5, "What Makes AgentGrid Innovative?")

    pillars = [
        ("1. OBJECTIVE → EXECUTION", "Natural-language business objectives translate directly into actionable multi-agent workflows.", ACCENT_CYAN),
        ("2. CEO ORCHESTRATION", "Master CEO Agent decomposes goals into dependency-aware Directed Acyclic Graphs (DAGs).", ACCENT_PURPLE),
        ("3. SPECIALIZED AI WORKFORCE", "Finance, Marketing, Hiring, & Legal operate as distinct, coordinated executive departments.", ACCENT_BLUE),
        ("4. HUMAN-IN-THE-LOOP", "Founder approval gates control plan authorization, deliverables, and external actions.", ACCENT_AMBER),
        ("5. CROSS-AGENT CONTEXT", "Agents share execution outputs and domain state instead of operating as isolated chatbots.", ACCENT_CYAN)
    ]

    for idx, (pil_title, pil_desc, pil_col) in enumerate(pillars):
        top_p = Inches(1.3 + idx * 0.98)
        add_card(s5, Inches(0.8), top_p, Inches(11.7), Inches(0.88), bg_color=CARD_BG, border_color=BORDER_DIM)
        
        tb_pil = s5.shapes.add_textbox(Inches(1.1), top_p + Inches(0.1), Inches(11.1), Inches(0.7))
        tf_pil = tb_pil.text_frame
        tf_pil.word_wrap = True
        
        pp1 = tf_pil.paragraphs[0]
        pp1.text = pil_title
        pp1.font.size = Pt(12)
        pp1.font.bold = True
        pp1.font.color.rgb = pil_col
        
        pp2 = tf_pil.add_paragraph()
        pp2.text = pil_desc
        pp2.font.size = Pt(10)
        pp2.font.color.rgb = TEXT_MUTED

    # Central Highlight Banner (Bottom)
    add_card(s5, Inches(0.8), Inches(6.3), Inches(11.7), Inches(0.7), bg_color=CARD_HEADER_BG, border_color=ACCENT_PURPLE)
    tb_inn_close = s5.shapes.add_textbox(Inches(1.0), Inches(6.42), Inches(11.3), Inches(0.5))
    pinc = tb_inn_close.text_frame.paragraphs[0]
    pinc.alignment = PP_ALIGN.CENTER
    pinc.text = '"From AI that answers questions → AI that coordinates business execution."'
    pinc.font.size = Pt(14)
    pinc.font.bold = True
    pinc.font.color.rgb = ACCENT_CYAN

    # =========================================================================
    # SLIDE 6 — IMPACT
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_background(s6)
    add_header(s6, "Impact")

    impact_sections = [
        ("FOUNDERS", [
            "Less repetitive operational work",
            "Faster decision-making",
            "Single executive command center",
            "Complete workflow visibility"
        ], ACCENT_CYAN, Inches(0.8), Inches(1.3)),

        ("STARTUPS", [
            "Automate repetitive operations",
            "Coordinate multiple departments",
            "Reduce operational overhead",
            "Scale execution efficiently"
        ], ACCENT_PURPLE, Inches(6.7), Inches(1.3)),

        ("RECRUITMENT", [
            "Structured candidate pipelines",
            "Consistent technical assessments",
            "Automated multi-round scoring",
            "Accelerated evaluation cycles"
        ], ACCENT_BLUE, Inches(0.8), Inches(4.1)),

        ("BUSINESS OPERATIONS", [
            "Connect disparate business functions",
            "Eliminate manual cross-dept overhead",
            "Turn objectives into measurable execution",
            "Maintain strict human governance"
        ], ACCENT_AMBER, Inches(6.7), Inches(4.1))
    ]

    for imp_title, imp_bullets, imp_col, left_pos, top_pos in impact_sections:
        add_card(s6, left_pos, top_pos, Inches(5.83), Inches(2.5), bg_color=CARD_BG, border_color=BORDER_DIM)
        
        tb_imp_hdr = s6.shapes.add_textbox(left_pos + Inches(0.3), top_pos + Inches(0.15), Inches(5.2), Inches(0.4))
        pih = tb_imp_hdr.text_frame.paragraphs[0]
        pih.text = imp_title
        pih.font.size = Pt(12)
        pih.font.bold = True
        pih.font.color.rgb = imp_col

        tb_imp_body = s6.shapes.add_textbox(left_pos + Inches(0.3), top_pos + Inches(0.55), Inches(5.2), Inches(1.8))
        tf_ib = tb_imp_body.text_frame
        tf_ib.word_wrap = True

        for idx, b_text in enumerate(imp_bullets):
            p = tf_ib.paragraphs[0] if idx == 0 else tf_ib.add_paragraph()
            p.text = f"✓  {b_text}"
            p.font.size = Pt(11)
            p.font.color.rgb = TEXT_WHITE

    # Bottom Impact Statement
    add_card(s6, Inches(0.8), Inches(6.7), Inches(11.7), Inches(0.5), bg_color=CARD_HEADER_BG, border_color=ACCENT_PURPLE)
    tb_imp_stmt = s6.shapes.add_textbox(Inches(1.0), Inches(6.75), Inches(11.3), Inches(0.4))
    pis = tb_imp_stmt.text_frame.paragraphs[0]
    pis.alignment = PP_ALIGN.CENTER
    pis.text = '"AgentGrid transforms fragmented business operations into one coordinated AI workflow."'
    pis.font.size = Pt(12)
    pis.font.bold = True
    pis.font.color.rgb = TEXT_MUTED

    # =========================================================================
    # SLIDE 7 — FEASIBILITY & FUTURE IMPACT
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_background(s7)
    add_header(s7, "Feasibility & Future Impact")

    # Left Container: Feasibility Today
    add_card(s7, Inches(0.8), Inches(1.3), Inches(5.7), Inches(4.8), bg_color=CARD_BG, border_color=ACCENT_CYAN)

    tb_f_hdr = s7.shapes.add_textbox(Inches(1.1), Inches(1.5), Inches(5.1), Inches(0.4))
    pfh = tb_f_hdr.text_frame.paragraphs[0]
    pfh.text = "FEASIBILITY TODAY"
    pfh.font.size = Pt(13)
    pfh.font.bold = True
    pfh.font.color.rgb = ACCENT_CYAN

    feasibility_points = [
        "Built with mature web & API standards (React 19, FastAPI).",
        "Google Gemini API provides the core AI reasoning layer.",
        "FastAPI powers production-grade async backend services.",
        "DAG architecture deterministically manages agent dependencies.",
        "Modular structure allows new agent domains to extend incrementally.",
        "Human approval gates provide absolute safety for consequential actions."
    ]

    tb_f_body = s7.shapes.add_textbox(Inches(1.1), Inches(2.0), Inches(5.1), Inches(3.9))
    tf_fb = tb_f_body.text_frame
    tf_fb.word_wrap = True

    for idx, fp in enumerate(feasibility_points):
        p = tf_fb.paragraphs[0] if idx == 0 else tf_fb.add_paragraph()
        p.text = f"• {fp}"
        p.font.size = Pt(10.5)
        p.font.color.rgb = TEXT_WHITE

    # Right Container: Future Roadmap
    add_card(s7, Inches(6.8), Inches(1.3), Inches(5.7), Inches(4.8), bg_color=CARD_BG, border_color=ACCENT_PURPLE)

    tb_r_hdr = s7.shapes.add_textbox(Inches(7.1), Inches(1.5), Inches(5.1), Inches(0.4))
    prh = tb_r_hdr.text_frame.paragraphs[0]
    prh.text = "FUTURE ROADMAP"
    prh.font.size = Pt(13)
    prh.font.bold = True
    prh.font.color.rgb = ACCENT_PURPLE

    roadmap_phases = [
        ("PHASE 1 — CURRENT", [
            "✓ CEO orchestration & DAG workflows",
            "✓ Finance, Marketing, Hiring & Legal agents",
            "✓ Human approval gate governance system",
            "✓ Candidate assessment & scoring workflow"
        ], ACCENT_CYAN),

        ("PHASE 2 — NEXT", [
            "→ Dedicated candidate interview platform",
            "→ Email / Calendar / Communication integrations",
            "→ Deep external business app integrations"
        ], ACCENT_BLUE),

        ("PHASE 3 — FUTURE", [
            "→ Additional specialized domain agents",
            "→ Deeper autonomous execution boundaries",
            "→ Scalable multi-company deployment architecture"
        ], ACCENT_PURPLE)
    ]

    for idx, (ph_name, ph_bullets, ph_col) in enumerate(roadmap_phases):
        top_rm = Inches(2.0 + idx * 1.3)
        add_card(s7, Inches(7.1), top_rm, Inches(5.1), Inches(1.2), bg_color=CARD_HEADER_BG, border_color=ph_col)
        
        tb_ph = s7.shapes.add_textbox(Inches(7.2), top_rm + Inches(0.08), Inches(4.9), Inches(1.0))
        tf_ph = tb_ph.text_frame
        tf_ph.word_wrap = True
        
        pph = tf_ph.paragraphs[0]
        pph.text = ph_name
        pph.font.size = Pt(10)
        pph.font.bold = True
        pph.font.color.rgb = ph_col

        for b_idx, b_txt in enumerate(ph_bullets):
            p = tf_ph.add_paragraph()
            p.text = b_txt
            p.font.size = Pt(9)
            p.font.color.rgb = TEXT_WHITE

    # Bottom Closing Box
    add_card(s7, Inches(0.8), Inches(6.25), Inches(11.7), Inches(0.9), bg_color=CARD_HEADER_BG, border_color=ACCENT_PURPLE)
    tb_c7 = s7.shapes.add_textbox(Inches(1.0), Inches(6.32), Inches(11.3), Inches(0.75))
    tf_c7 = tb_c7.text_frame
    tf_c7.word_wrap = True

    pc7_1 = tf_c7.paragraphs[0]
    pc7_1.alignment = PP_ALIGN.CENTER
    pc7_1.text = '"AI can execute the work. Humans remain in control of the decisions."'
    pc7_1.font.size = Pt(13)
    pc7_1.font.bold = True
    pc7_1.font.color.rgb = TEXT_WHITE

    pc7_2 = tf_c7.add_paragraph()
    pc7_2.alignment = PP_ALIGN.CENTER
    pc7_2.text = 'AgentGrid — "Where AI Agents Work Together."'
    pc7_2.font.size = Pt(11)
    pc7_2.font.bold = True
    pc7_2.font.color.rgb = ACCENT_CYAN

    output_path = "c:\\Users\\Yogendar\\Downloads\\AgentGrid\\AgentGrid_Hackathon_Presentation.pptx"
    prs.save(output_path)
    print(f"SUCCESSFULLY GENERATED PRESENTATION AT: {output_path}")

if __name__ == "__main__":
    build_presentation()
